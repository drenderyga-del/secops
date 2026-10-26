from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.asoc_products import (
    ASOC_MATCH_CONFLICT,
    ASOC_MATCH_MAPPED,
    ASOC_MATCH_MISSING_TARGET,
    ASOC_MATCH_NO_REPOSITORIES,
    ASOC_MATCH_PARTIAL,
    AsocSyncConfirmationRequiredError,
    read_sync_outcome,
    run_scheduled_sync,
)
from app.models import (
    ActivePipeline,
    AsocProduct,
    AsocProductRepositoryMapping,
    AuditCoverageStatus,
    AuditRepository,
    AuditSnapshot,
    JobReport,
    Product,
    Repository,
)
from app.models.database import Base
from app.models.job_report import JobStatus
from app.models.pipeline import PipelineStatus
from app.audit_service import AuditService
from app.product_registry import (
    EXPORT_COLUMNS,
    Quarter,
    RegistryQuery,
    build_registry,
    export_filename,
    export_xlsx,
    product_detail,
)

NOW = datetime(2026, 9, 26, 12, 0, tzinfo=timezone.utc)


@pytest.fixture(autouse=True)
def configured_gitlab_hosts(monkeypatch):
    monkeypatch.setenv("ASOC_PRODUCT_REFERENCE_HOSTS", "product.example.com")
    monkeypatch.setattr("app.gitlab_url.known_gitlab_hosts", lambda: {"gitlab.example"})


def _db():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _card(asoc_id, created, *, state=ASOC_MATCH_MAPPED, active=True, removed_at=None, description="", **extra):
    return AsocProduct(
        asoc_id=asoc_id,
        name=f"{asoc_id.title()} (BI_{extra.pop('bi', '1')})",
        title=asoc_id.title(),
        bi_id=extra.pop("bi_id", "1"),
        description=description,
        asoc_created=created,
        mapping_state=state,
        is_active=active,
        removed_at=removed_at,
        **extra,
    )


def _keys(registry):
    return [row["asoc_id"] for row in registry["rows"]]


def test_quarter_boundaries_use_the_operator_timezone():
    quarter = Quarter.parse("2026-q3")

    assert quarter.key == "2026-Q3"
    assert quarter.label == "2026 Q3"
    # 1 July 00:00 in Moscow is 30 June 21:00 UTC.
    assert quarter.start == datetime(2026, 6, 30, 21, 0, tzinfo=timezone.utc)
    assert quarter.end == datetime(2026, 9, 30, 21, 0, tzinfo=timezone.utc)
    assert Quarter.containing(datetime(2026, 6, 30, 22, 30, tzinfo=timezone.utc)).key == "2026-Q3"


@pytest.mark.parametrize("value", ["2026", "2026-Q5", "Q3-2026", "1999-Q1", "2026-Q"])
def test_invalid_periods_are_rejected(value):
    with pytest.raises(ValueError):
        RegistryQuery.from_params(period=value)


def test_quarter_slice_new_and_total_include_cards_removed_later():
    db = _db()
    db.add_all([
        _card("q1", datetime(2026, 2, 1, tzinfo=timezone.utc)),
        _card("q2", datetime(2026, 5, 1, tzinfo=timezone.utc)),
        # Created in Q1 and removed from ASOC during Q3.
        _card("gone", datetime(2026, 1, 15, tzinfo=timezone.utc), active=False,
              removed_at=datetime(2026, 8, 1, tzinfo=timezone.utc)),
        # Midnight in Moscow on 1 July belongs to Q3 even though it is 30 June in UTC.
        _card("q3-edge", datetime(2026, 6, 30, 21, 30, tzinfo=timezone.utc)),
        _card("undated", None),
    ])
    db.commit()

    new_q1 = build_registry(db, RegistryQuery.from_params(period="2026-Q1", mode="new"), now=NOW)
    total_q2 = build_registry(db, RegistryQuery.from_params(period="2026-Q2", mode="total"), now=NOW)
    total_q3 = build_registry(db, RegistryQuery.from_params(period="2026-Q3", mode="total"), now=NOW)
    new_q3 = build_registry(db, RegistryQuery.from_params(period="2026-Q3", mode="new"), now=NOW)
    current = build_registry(db, RegistryQuery(), now=NOW)

    assert sorted(_keys(new_q1)) == ["gone", "q1"]
    assert next(row for row in new_q1["rows"] if row["asoc_id"] == "gone")["status"] == "removed"
    assert sorted(_keys(total_q2)) == ["gone", "q1", "q2"]
    assert total_q2["summary"]["new_in_period"] == 1
    assert sorted(_keys(total_q3)) == ["q1", "q2", "q3-edge"]
    assert _keys(new_q3) == ["q3-edge"]
    assert new_q3["rows"][0]["connected_quarter"] == "2026 Q3"
    assert sorted(_keys(current)) == ["q1", "q2", "q3-edge", "undated"]
    assert current["summary"]["undated"] == 1
    assert [period["key"] for period in current["periods"]] == ["2026-Q3", "2026-Q2", "2026-Q1"]


def test_statuses_and_summary_counts_ignore_the_status_filter():
    db = _db()
    created = datetime(2026, 9, 1, tzinfo=timezone.utc)
    db.add_all([
        _card("synced", created),
        _card("no-link", created, state=ASOC_MATCH_NO_REPOSITORIES),
        _card("waiting", NOW - timedelta(days=10), state=ASOC_MATCH_MISSING_TARGET,
              description="https://gitlab.example/new/group"),
        _card("fresh", NOW - timedelta(days=1), state=ASOC_MATCH_MISSING_TARGET),
        _card("conflict", created, state=ASOC_MATCH_CONFLICT),
    ])
    db.commit()

    registry = build_registry(db, RegistryQuery.from_params(status="attention"), now=NOW)
    by_id = {row["asoc_id"]: row for row in registry["rows"]}

    assert set(by_id) == {"no-link", "waiting", "fresh", "conflict"}
    assert registry["summary"]["total"] == 5
    assert registry["summary"]["synced"] == 1
    assert registry["summary"]["attention"] == 4
    assert by_id["no-link"]["status"] == "no_gitlab_link"
    assert by_id["conflict"]["status"] == "conflict"
    assert by_id["waiting"]["status"] == "group_not_monitored"
    assert by_id["waiting"]["waiting_days"] == 10
    assert by_id["waiting"]["overdue"] is True
    assert by_id["fresh"]["overdue"] is False
    # no-link and conflict were connected 25 days ago, waiting 10 days ago;
    # fresh (1 day) is still inside the grace period.
    assert registry["summary"]["overdue"] == 3


def test_scope_filter_and_search_by_bi_id_or_group():
    db = _db()
    group = Product(name="example_app", path="sandbox/example_app", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add_all([
        _card("internal", NOW, description="ScanOps-Scope: internal", bi_id="241"),
        _card("external", NOW, description="ScanOps-Scope: external", bi_id="242", monitor_product_id=group.id),
    ])
    db.commit()

    assert _keys(build_registry(db, RegistryQuery.from_params(scope="internal"), now=NOW)) == ["internal"]
    assert _keys(build_registry(db, RegistryQuery.from_params(q="bi_242"), now=NOW)) == ["external"]
    assert _keys(build_registry(db, RegistryQuery.from_params(q="sandbox/example"), now=NOW)) == ["external"]


def test_card_links_open_only_mapped_groups_never_bi_ids():
    db = _db()
    db.add_all([
        Product(id=3, name="example_app", path="sandbox/example_app", gitlab_base_url="https://gitlab.example"),
    ])
    db.flush()
    db.add_all([
        _card("mapped", NOW, bi_id="305", monitor_product_id=3),
        _card("orphan", NOW, bi_id="404", state=ASOC_MATCH_MISSING_TARGET),
    ])
    db.commit()

    by_id = {row["asoc_id"]: row for row in build_registry(db, RegistryQuery.from_params(), now=NOW)["rows"]}

    assert by_id["mapped"]["href"] == "/targets/products/asoc/mapped"
    assert [group["href"] for group in by_id["mapped"]["groups"]] == ["/targets/products/3"]
    assert by_id["orphan"]["href"] == "/targets/products/asoc/orphan"
    assert by_id["orphan"]["groups"] == []
    hrefs = repr(by_id)
    assert "/targets/products/305" not in hrefs
    assert "/targets/products/404" not in hrefs


def test_sibling_groups_in_one_namespace_keep_their_own_audited_repositories():
    db = _db()
    first = Product(name="one", path="qa/product-01", gitlab_base_url="https://gitlab.example")
    second = Product(name="two", path="qa/product-02", gitlab_base_url="https://gitlab.example")
    db.add_all([first, second])
    db.flush()
    db.add(Repository(product_id=first.id, path="qa/product-01/api"))
    # Audit-only projects: never reported to ScanOps, attributed by group path.
    for index, path in enumerate(
        ["qa/product-01/api", "qa/product-01/unscanned", "qa/product-02/unscanned", "qa/other/tool"],
        start=1,
    ):
        repo = AuditRepository(
            instance_host="gitlab.example",
            project_id=index,
            path_with_namespace=path,
            web_url=f"https://gitlab.example/{path}",
        )
        db.add(repo)
        db.flush()
        db.add(AuditSnapshot(
            audit_repository_id=repo.id,
            scanned_at=NOW - timedelta(hours=1),
            status=AuditCoverageStatus.PIPELINE_NO_SEC,
        ))
    db.commit()

    def paths(product):
        return sorted(item["path"] for item in product_detail(db, product_id=product.id, now=NOW)["repositories"])

    assert paths(first) == ["qa/product-01/api", "qa/product-01/unscanned"]
    assert paths(second) == ["qa/product-02/unscanned"]
    # Coverage uses the same attribution; only the project outside every
    # product group falls back to the shared namespace.
    coverage = {
        product_id: sorted(repo.path_with_namespace for _snapshot, repo in rows)
        for product_id, rows in AuditService(db).latest_repositories_by_product().items()
    }
    assert coverage[second.id] == ["qa/product-02/unscanned"]
    assert "qa/product-02/unscanned" not in coverage[first.id]


@pytest.mark.parametrize("audited", [False, True])
def test_same_project_path_on_distinct_gitlab_instances_keeps_separate_evidence(audited):
    db = _db()
    card = _card("multi-instance", NOW - timedelta(days=1))
    db.add(card)
    expected = {}
    for index, host in enumerate(("gitlab-a.example", "gitlab-b.example"), start=1):
        group = Product(name=f"Group {index}", path="team/shared", gitlab_base_url=f"https://{host}")
        db.add(group)
        db.flush()
        repository = Repository(product_id=group.id, path="team/shared/api")
        db.add(repository)
        db.flush()
        url = f"https://{host}/team/shared/api"
        expected[repository.id] = (group.id, url)
        db.add(AsocProductRepositoryMapping(
            asoc_id=card.asoc_id,
            product_id=group.id,
            description_url=f"https://{host}/team/shared",
        ))
        db.add(ActivePipeline(
            project_id=index,
            pipeline_id=100 + index,
            repository_id=repository.id,
            product_id=group.id,
            status=PipelineStatus.COMPLETED if index == 1 else PipelineStatus.FAILED,
            completed_at=NOW - timedelta(hours=1),
            web_url=f"{url}/-/pipelines/{100 + index}",
        ))
        if audited:
            audit = AuditRepository(
                instance_host=host,
                project_id=index,
                path_with_namespace=repository.path,
                web_url=url,
            )
            db.add(audit)
            db.flush()
            db.add(AuditSnapshot(
                audit_repository_id=audit.id,
                scanned_at=NOW - timedelta(minutes=30),
                status=AuditCoverageStatus.COVERED if index == 1 else AuditCoverageStatus.PIPELINE_NO_SEC,
            ))
    db.commit()

    registry = build_registry(db, RegistryQuery(), now=NOW)
    row = registry["rows"][0]
    assert row["repositories"]["total"] == 2
    assert row["repositories"]["monitored"] == 2
    assert row["repositories"]["scanned"] == (1 if audited else 0)
    assert row["pipelines"]["completed"] == 1
    assert row["pipelines"]["failed"] == 1
    assert row["pipelines"]["success_rate"] == 50
    assert registry["summary"]["repositories"]["total"] == 2
    detail = product_detail(db, asoc_id=card.asoc_id, now=NOW)
    assert {
        item["scanops_repository_id"]: (item["group_id"], item["url"])
        for item in detail["repositories"]
    } == expected
    for item in detail["repositories"]:
        assert item["jobs_href"] == (
            f"/repos/{item['scanops_repository_id']}/jobs?only_failed=true&status=failed&timeframe=all"
        )
    workbook = load_workbook(BytesIO(export_xlsx(registry)))
    summary = dict(workbook["Summary"].iter_rows(values_only=True))
    assert summary["Repositories in groups"] == 2
    assert summary["Pipeline success, 30d"] == 50
    db.close()


def _product_with_evidence(db):
    group = Product(name="example_app", path="sandbox/example_app", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    api = Repository(product_id=group.id, path="sandbox/example_app/api")
    web = Repository(product_id=group.id, path="sandbox/example_app/web")
    # Known to ScanOps but not yet observed by the GitLab audit.
    worker = Repository(product_id=group.id, path="sandbox/example_app/worker")
    db.add_all([api, web, worker])
    db.flush()
    audit = {
        "sandbox/example_app/api": AuditCoverageStatus.COVERED,
        "sandbox/example_app/web": AuditCoverageStatus.COVERED,
        "sandbox/example_app/batch": AuditCoverageStatus.PIPELINE_NO_SEC,
        "sandbox/example_app/legacy": AuditCoverageStatus.NO_PIPELINES,
        "sandbox/example_app/old": AuditCoverageStatus.EXCLUDED,
    }
    for index, (path, status) in enumerate(audit.items()):
        repo = AuditRepository(
            instance_host="gitlab.example",
            project_id=index + 1,
            path_with_namespace=path,
            web_url=f"https://gitlab.example/{path}",
            excluded=status == AuditCoverageStatus.EXCLUDED,
        )
        db.add(repo)
        db.flush()
        db.add(AuditSnapshot(audit_repository_id=repo.id, scanned_at=NOW - timedelta(hours=2), status=status))
    pipelines = [
        (api, PipelineStatus.COMPLETED, NOW - timedelta(days=2)),
        (api, PipelineStatus.COMPLETED, NOW - timedelta(days=1)),
        (web, PipelineStatus.COMPLETED, NOW - timedelta(days=3)),
        (web, PipelineStatus.FAILED, NOW - timedelta(hours=5)),
        # Outside the 30-day window: counts for "last" but not for the rate.
        (worker, PipelineStatus.FAILED, NOW - timedelta(days=45)),
    ]
    for index, (repository, status, at) in enumerate(pipelines):
        db.add(ActivePipeline(
            project_id=repository.id,
            pipeline_id=100 + index,
            repository_id=repository.id,
            product_id=group.id,
            status=status,
            completed_at=at,
            web_url=f"https://gitlab.example/{repository.path}/-/pipelines/{100 + index}",
        ))
    db.add(JobReport(
        repository_id=web.id,
        job_name="sca_scan",
        job_status=JobStatus.FAILED,
        ref="main",
        finished_at=NOW - timedelta(hours=5),
    ))
    card = _card(
        "example_app",
        datetime(2026, 7, 3, tzinfo=timezone.utc),
        bi_id="241",
        description=(
            "ScanOps-Scope: internal\n"
            "https://gitlab.example/sandbox/example_app\n"
            "https://gitlab.example/sandbox/widgets\n"
            "https://product.example.com/products/241"
        ),
        monitor_product_id=group.id,
    )
    db.add(card)
    db.flush()
    db.add(AsocProductRepositoryMapping(
        asoc_id="example_app",
        product_id=group.id,
        description_url="https://gitlab.example/sandbox/example_app",
    ))
    db.commit()
    return group, api, web, worker


def test_repository_coverage_and_pipeline_success_come_from_group_evidence():
    db = _db()
    _product_with_evidence(db)

    row = build_registry(db, RegistryQuery(), now=NOW)["rows"][0]

    assert row["repositories"] == {
        "total": 5,
        "scanned": 2,
        "not_scanned": 1,
        "inactive": 1,
        "unknown": 0,
        "not_audited": 1,
        "excluded": 1,
        "monitored": 3,
        "audited": True,
        "audited_at": (NOW - timedelta(hours=2)).isoformat(),
    }
    assert row["pipelines"]["completed"] == 3
    assert row["pipelines"]["failed"] == 1
    assert row["pipelines"]["success_rate"] == 75
    assert row["pipelines"]["repos_with_pipelines"] == 3
    assert row["pipelines"]["repos_green"] == 1
    assert row["links_total"] == 2
    assert row["links_matched"] == 1


def test_product_detail_explains_links_and_orders_repositories_by_attention():
    db = _db()
    group, _api, web, _worker = _product_with_evidence(db)

    detail = product_detail(db, asoc_id="example_app", now=NOW)

    assert detail["kind"] == "asoc"
    assert [(link["url"], link["state"]) for link in detail["links"]] == [
        ("https://gitlab.example/sandbox/example_app", "matched"),
        ("https://gitlab.example/sandbox/widgets", "not_monitored"),
        ("https://product.example.com/products/241", "reference"),
    ]
    statuses = [repo["status"] for repo in detail["repositories"]]
    assert statuses == ["not_scanned", "not_audited", "inactive", "scanned", "scanned", "excluded"]
    web_row = next(repo for repo in detail["repositories"] if repo["path"] == "sandbox/example_app/web")
    assert web_row["open_failed_jobs"] == 1
    assert web_row["pipelines"]["last_status"] == "failed"
    assert web_row["href"] == f"/targets/repositories/{web.id}"
    assert detail["open_failed_jobs"] == 1

    via_group = product_detail(db, product_id=group.id, now=NOW)
    assert via_group["kind"] == "asoc"
    assert via_group["product"]["asoc_id"] == "example_app"
    assert via_group["focus_group_id"] == group.id


def test_group_without_asoc_card_is_listed_and_shows_claiming_cards():
    db = _db()
    # What a real sync leaves behind when cards compete for a group.
    group = Product(name="orphan", path="sandbox/orphan", gitlab_base_url="https://gitlab.example",
                    asoc_sync_status=ASOC_MATCH_CONFLICT)
    db.add(group)
    db.flush()
    db.add(_card("claim", NOW, state=ASOC_MATCH_CONFLICT, description="https://gitlab.example/sandbox/orphan"))
    db.commit()

    registry = build_registry(db, RegistryQuery(), now=NOW)
    detail = product_detail(db, product_id=group.id, now=NOW)

    assert [(row["id"], row["asoc_conflict"]) for row in registry["unregistered_groups"]] == [(group.id, True)]
    assert detail["kind"] == "group"
    assert [(card["asoc_id"], card["state"]) for card in detail["claimed_by"]] == [("claim", "conflict")]
    assert product_detail(db, asoc_id="missing", now=NOW) is None

    # A card that only links the group (not yet synced) is not a conflict.
    db.query(AsocProduct).filter_by(asoc_id="claim").one().mapping_state = ASOC_MATCH_MISSING_TARGET
    db.commit()
    detail = product_detail(db, product_id=group.id, now=NOW)
    assert [(card["asoc_id"], card["state"]) for card in detail["claimed_by"]] == [("claim", "pending_sync")]


def test_export_matches_the_view_and_keeps_asoc_text_inert():
    db = _db()
    _product_with_evidence(db)
    db.add(_card("formula", datetime(2026, 8, 1, tzinfo=timezone.utc), state=ASOC_MATCH_NO_REPOSITORIES))
    db.query(AsocProduct).filter_by(asoc_id="formula").one().title = "=HYPERLINK(\"x\")"
    db.commit()

    registry = build_registry(db, RegistryQuery.from_params(period="2026-Q3", mode="new"), now=NOW)
    workbook = load_workbook(BytesIO(export_xlsx(registry)))
    products = workbook["Products"]
    rows = list(products.iter_rows(values_only=True))

    assert list(rows[0]) == [name for name, _width in EXPORT_COLUMNS]
    assert len(rows) == 1 + len(registry["rows"])
    example_app = next(row for row in rows[1:] if row[1] == "Example_App")
    assert example_app[0] == "BI_241"
    assert example_app[3] == "sandbox/example_app"
    assert example_app[5:11] == (5, 2, 1, 1, 75, "1/3")
    assert example_app[12] == "2026-07-03"
    assert example_app[13] == "2026 Q3"
    formula = next(row for row in rows[1:] if row[1] and row[1].startswith("=HYPERLINK"))
    assert products.cell(row=rows.index(formula) + 1, column=2).data_type == "s"
    summary = dict(workbook["Summary"].iter_rows(values_only=True))
    assert summary["Period"] == "2026 Q3 · new in quarter"
    assert summary["New in quarter"] == 2
    assert export_filename(registry) == "scanops-products-2026-Q3-new.xlsx"


def test_summary_counts_a_group_once_when_a_removed_card_was_replaced():
    db = _db()
    group = Product(name="example_app", path="sandbox/example_app", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add_all([
        Repository(product_id=group.id, path="sandbox/example_app/api"),
        Repository(product_id=group.id, path="sandbox/example_app/web"),
    ])
    created = datetime(2026, 7, 10, tzinfo=timezone.utc)
    db.add_all([
        _card("old-card", created, monitor_product_id=group.id, active=False, removed_at=created + timedelta(days=20),
              state=ASOC_MATCH_MAPPED),
        _card("new-card", created + timedelta(days=21), monitor_product_id=group.id),
    ])
    db.commit()

    registry = build_registry(db, RegistryQuery.from_params(period="2026-Q3", mode="new"), now=NOW)

    assert {row["asoc_id"] for row in registry["rows"]} == {"old-card", "new-card"}
    assert all(row["repositories"]["total"] == 2 for row in registry["rows"])
    assert registry["summary"]["repositories"]["total"] == 2


def test_export_drops_control_characters_instead_of_failing():
    db = _db()
    db.add(_card("ctrl", datetime(2026, 8, 1, tzinfo=timezone.utc), state=ASOC_MATCH_NO_REPOSITORIES))
    db.query(AsocProduct).filter_by(asoc_id="ctrl").one().title = "Vertical\x0btab\x01product"
    db.commit()

    registry = build_registry(db, RegistryQuery(), now=NOW)
    workbook = load_workbook(BytesIO(export_xlsx(registry)))

    names = [row[1] for row in workbook["Products"].iter_rows(min_row=2, values_only=True)]
    assert names == ["Verticaltabproduct"]


def test_card_links_encode_asoc_controlled_ids_and_open_them(monkeypatch):
    monkeypatch.delenv("UI_AUTH_USERNAME", raising=False)
    monkeypatch.delenv("UI_AUTH_PASSWORD", raising=False)
    from sqlalchemy.pool import StaticPool

    from app.main import app
    from app.routers.web.common import get_db

    # Sync FastAPI routes run in a worker thread.
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(_card("x?y#z/1", NOW, state=ASOC_MATCH_NO_REPOSITORIES))
    db.commit()

    row = build_registry(db, RegistryQuery(), now=NOW)["rows"][0]
    assert row["href"] == "/targets/products/asoc/x%3Fy%23z%2F1"

    app.dependency_overrides[get_db] = lambda: db
    try:
        client = TestClient(app)
        page = client.get(row["href"])
        response = client.get("/api/ui/products/widgets/x%3Fy%23z%2F1")
    finally:
        app.dependency_overrides.pop(get_db, None)
    assert page.status_code == 200
    assert response.status_code == 200, response.text
    assert response.json()["product"]["asoc_id"] == "x?y#z/1"


def test_attribution_breaks_ties_by_name_and_scopes_pathless_products_by_full_name():
    from app.audit_service import RepositoryAttribution

    db = _db()
    db.add_all([
        Product(name="zeta", path="qa/shared", gitlab_base_url="https://gitlab.example"),
        Product(name="alpha", path="qa/shared", gitlab_base_url="https://gitlab.example"),
        Product(name="team/app", path="", gitlab_base_url="https://gitlab.example"),
    ])
    db.commit()
    attribution = RepositoryAttribution(db)

    for _ in range(3):
        assert attribution.product_for("gitlab.example", "qa/shared/api").name == "alpha"
    assert attribution.product_for("gitlab.example", "team/app/service").name == "team/app"
    # The pathless product no longer claims its whole top-level namespace.
    assert attribution.product_for("gitlab.example", "team/other/service") is None


def test_out_of_range_asoc_dates_are_undated_instead_of_breaking_the_registry():
    from app.asoc_products import _parse_created

    assert _parse_created("9999-12-31T23:00:00Z") is None
    assert _parse_created("0001-01-01T00:00:00Z") is None
    assert _parse_created("2026-07-03T10:00:00Z").year == 2026

    db = _db()
    db.add_all([
        _card("future", datetime(9999, 12, 31, 23, tzinfo=timezone.utc), state=ASOC_MATCH_NO_REPOSITORIES),
        _card("ancient", datetime(1, 1, 1, tzinfo=timezone.utc), state=ASOC_MATCH_NO_REPOSITORIES),
        _card("normal", datetime(2026, 8, 1, tzinfo=timezone.utc), state=ASOC_MATCH_NO_REPOSITORIES),
    ])
    db.commit()

    registry = build_registry(db, RegistryQuery(), now=NOW)
    quarter = build_registry(db, RegistryQuery.from_params(period="2026-Q3", mode="total"), now=NOW)

    assert registry["summary"]["undated"] == 2
    assert [period["key"] for period in registry["periods"]] == ["2026-Q3"]
    assert _keys(quarter) == ["normal"]
    assert export_xlsx(registry)


@pytest.mark.parametrize(
    ("description", "expected"),
    [
        ("https://gitlab.example/sandbox/orphan", "pending_sync"),
        ("https://gitlab.example/sandbox/orphan/api", "pending_sync"),
        ("https://gitlab.example/sandbox/orphan/audit-only", "unresolved"),
    ],
)
def test_group_page_tells_pending_cards_from_unresolvable_links(description, expected):
    db = _db()
    group = Product(name="orphan", path="sandbox/orphan", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add(Repository(product_id=group.id, path="sandbox/orphan/api"))
    repo = AuditRepository(
        instance_host="gitlab.example",
        project_id=1,
        path_with_namespace="sandbox/orphan/audit-only",
        web_url="https://gitlab.example/sandbox/orphan/audit-only",
    )
    db.add(repo)
    db.flush()
    db.add(AuditSnapshot(audit_repository_id=repo.id, scanned_at=NOW, status=AuditCoverageStatus.PIPELINE_NO_SEC))
    db.add(_card("claim", NOW, state=ASOC_MATCH_MISSING_TARGET, description=description))
    db.commit()

    [claim] = product_detail(db, product_id=group.id, now=NOW)["claimed_by"]

    assert claim["state"] == expected
    assert claim["status_label"] == "Group not in ScanOps"


def test_group_claim_matches_what_the_sync_can_actually_resolve():
    db = _db()
    # A group path stored as a full URL is displayed fine but never resolved
    # by the sync, so a card linking it must not be promised a connection.
    group = Product(name="url-path", path="https://gitlab.example/sandbox/url-path", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add(_card("links-group", NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/url-path"))
    db.commit()

    claims = product_detail(db, product_id=group.id, now=NOW)["claimed_by"]

    assert [(claim["asoc_id"], claim["state"]) for claim in claims] == [("links-group", "unresolved")]


def test_two_cards_waiting_for_one_group_are_shown_as_a_conflict():
    db = _db()
    group = Product(name="shared", path="sandbox/shared", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add_all([
        _card(asoc_id, NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/shared")
        for asoc_id in ("first", "second")
    ])
    db.commit()

    claims = product_detail(db, product_id=group.id, now=NOW)["claimed_by"]

    assert {claim["state"] for claim in claims} == {"conflict"}


def test_group_url_shared_by_two_products_blocks_only_cards_linking_the_group():
    db = _db()
    first = Product(name="a-shared", path="sandbox/shared", gitlab_base_url="https://gitlab.example")
    second = Product(name="b-shared", path="sandbox/shared", gitlab_base_url="https://gitlab.example")
    db.add_all([first, second])
    db.flush()
    db.add(Repository(product_id=first.id, path="sandbox/shared/api"))
    db.add(_card("by-group", NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/shared"))
    db.commit()

    [claim] = product_detail(db, product_id=first.id, now=NOW)["claimed_by"]
    assert claim["state"] == "conflict"

    db.query(AsocProduct).filter_by(asoc_id="by-group").one().description = "https://gitlab.example/sandbox/shared/api"
    db.commit()
    [claim] = product_detail(db, product_id=first.id, now=NOW)["claimed_by"]
    assert claim["state"] == "pending_sync"


def test_card_on_a_shared_group_url_does_not_block_a_unique_repository_link():
    db = _db()
    first = Product(name="a-shared", path="sandbox/shared", gitlab_base_url="https://gitlab.example")
    second = Product(name="b-shared", path="sandbox/shared", gitlab_base_url="https://gitlab.example")
    db.add_all([first, second])
    db.flush()
    db.add(Repository(product_id=first.id, path="sandbox/shared/api"))
    db.add_all([
        _card("by-group", NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/shared"),
        _card("by-repo", NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/shared/api"),
    ])
    db.commit()

    claims = {claim["asoc_id"]: claim["state"] for claim in product_detail(db, product_id=first.id, now=NOW)["claimed_by"]}

    assert claims == {"by-group": "conflict", "by-repo": "pending_sync"}


def test_a_card_linked_once_is_not_reported_as_waiting_when_it_degrades():
    db = _db()
    created = datetime(2024, 1, 1, tzinfo=timezone.utc)
    db.add_all([
        _card("degraded", created, state=ASOC_MATCH_PARTIAL, first_connected_at=created + timedelta(days=1)),
        _card("never-linked", created, state=ASOC_MATCH_MISSING_TARGET),
    ])
    db.commit()

    rows = {row["asoc_id"]: row for row in build_registry(db, RegistryQuery(), now=NOW)["rows"]}

    assert (rows["degraded"]["waiting_days"], rows["degraded"]["overdue"]) == (None, False)
    assert rows["never-linked"]["overdue"] is True


def test_card_linked_before_date_tracking_is_not_waiting():
    db = _db()
    group = Product(name="legacy", path="sandbox/legacy", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    created = datetime(2024, 1, 1, tzinfo=timezone.utc)
    db.add(_card("legacy", created, state=ASOC_MATCH_PARTIAL, monitor_product_id=group.id))
    db.commit()

    row = build_registry(db, RegistryQuery(), now=NOW)["rows"][0]

    assert (row["waiting_days"], row["overdue"], row["linked_before_tracking"]) == (None, False, True)
    assert build_registry(db, RegistryQuery(), now=NOW)["summary"]["overdue"] == 0


def test_pre_tracking_card_that_lost_its_links_is_still_not_waiting():
    db = _db()
    created = datetime(2024, 1, 1, tzinfo=timezone.utc)
    db.add(_card("lost", created, state=ASOC_MATCH_MISSING_TARGET, connected_before_tracking=True))
    db.commit()

    row = build_registry(db, RegistryQuery(), now=NOW)["rows"][0]

    assert (row["waiting_days"], row["overdue"], row["linked_before_tracking"]) == (None, False, True)


def test_new_card_joining_a_conflicted_group_is_not_promised_a_connection():
    db = _db()
    group = Product(name="contested", path="sandbox/contested", gitlab_base_url="https://gitlab.example")
    db.add(group)
    db.flush()
    db.add_all([
        _card("old", NOW, state=ASOC_MATCH_CONFLICT, description="https://gitlab.example/sandbox/contested"),
        _card("new", NOW, state=ASOC_MATCH_MISSING_TARGET, description="https://gitlab.example/sandbox/contested"),
    ])
    db.commit()

    claims = {claim["asoc_id"]: claim["state"] for claim in product_detail(db, product_id=group.id, now=NOW)["claimed_by"]}

    assert claims == {"old": "conflict", "new": "conflict"}


def test_undated_count_follows_audience_and_search_filters():
    db = _db()
    db.add_all([
        _card("undated-internal", None, description="ScanOps-Scope: internal", bi_id="301"),
        _card("undated-external", None, description="ScanOps-Scope: external", bi_id="302"),
        _card("undated-removed", None, active=False, removed_at=NOW, bi_id="303"),
    ])
    db.commit()

    assert build_registry(db, RegistryQuery(), now=NOW)["summary"]["undated"] == 2
    assert build_registry(db, RegistryQuery.from_params(scope="internal"), now=NOW)["summary"]["undated"] == 1
    assert build_registry(db, RegistryQuery.from_params(q="BI_302"), now=NOW)["summary"]["undated"] == 1
    quarter = RegistryQuery.from_params(period="2026-Q3", mode="total")
    assert build_registry(db, quarter, now=NOW)["summary"]["undated"] == 3


class _EnabledClient:
    def enabled(self):
        return True


def test_scheduled_sync_never_confirms_mass_removal(monkeypatch):
    db = _db()
    monkeypatch.setattr("app.asoc_products.AsocClient", _EnabledClient)

    def refuse(*_args, **kwargs):
        assert kwargs["confirmed_by"] == "scheduler"
        assert kwargs["unattended"] is True
        assert "confirm_mass_removal" not in kwargs
        raise AsocSyncConfirmationRequiredError(
            active_count=8, incoming_count=0, removal_count=8, registry_version="0" * 64,
        )

    monkeypatch.setattr("app.asoc_products.fetch_and_sync", refuse)
    outcome = run_scheduled_sync(db)

    assert outcome["status"] == "needs_confirmation"
    assert "8 would be removed" in outcome["message"]
    assert read_sync_outcome(db)["trigger"] == "scheduler"


def test_scheduled_sync_records_success_and_skips_when_not_configured(monkeypatch):
    db = _db()
    monkeypatch.setattr("app.asoc_products.AsocClient", _EnabledClient)
    monkeypatch.setattr(
        "app.asoc_products.fetch_and_sync",
        lambda *_args, **_kwargs: {"total": 3, "mapped": 2, "partial": 0, "missing_target": 1, "removed": 0},
    )

    outcome = run_scheduled_sync(db)

    assert outcome["status"] == "ok"
    assert outcome["message"] == "3 ASOC products: 2 synced with ScanOps, 1 need attention, 0 removed."
    assert outcome["result"]["mapped"] == 2

    class _Disabled:
        def enabled(self):
            return False

    monkeypatch.setattr("app.asoc_products.AsocClient", _Disabled)
    assert run_scheduled_sync(db) is None


def test_retired_report_routes_redirect_to_the_products_registry(monkeypatch):
    monkeypatch.delenv("UI_AUTH_USERNAME", raising=False)
    monkeypatch.delenv("UI_AUTH_PASSWORD", raising=False)
    from app.main import app

    client = TestClient(app)
    for garbage in ("year=1999&quarter=1", "year=%C2%B2&quarter=1", f"year={'9' * 5000}&quarter=1", "year=2026&quarter=5"):
        for path in ("/reports", "/reports/operator", "/reports/export.xlsx", "/classic/reports"):
            response = client.get(f"{path}?{garbage}", follow_redirects=False)
            assert response.status_code == 308, (path, garbage[:20])
            assert "period=" not in response.headers["location"]
    alias = client.get("/next?view=reports&year=2026&quarter=3&no_bi=true", follow_redirects=False)
    assert alias.headers["location"] == "/targets?level=product&period=2026-Q3&mode=new"
    report = client.get("/reports?year=2026&quarter=3", follow_redirects=False)
    operator = client.get("/reports/operator", follow_redirects=False)
    export = client.get("/reports/export.xlsx?year=2026&quarter=2&no_bi=true", follow_redirects=False)

    assert report.status_code == 308
    assert report.headers["location"] == "/targets?level=product&period=2026-Q3&mode=new"
    assert operator.headers["location"] == "/targets?level=product"
    assert export.headers["location"] == "/api/ui/products/export.xlsx?period=2026-Q2&mode=new"
