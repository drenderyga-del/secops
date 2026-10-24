"""Product registry: ASOC product cards joined with ScanOps evidence.

ASOC is the system of record for products. A card that appears in ASOC is a
connected product; its ASOC ``created`` timestamp is the connection date used
for quarterly reporting. After that ScanOps must *sync* the card: resolve the
GitLab links from the card description to monitored product groups, and then
observe security scans in the repositories of those groups.

The same projection feeds the Products table, the per-product drill-down and
the XLSX export, so an exported quarter always equals what the operator sees.
"""
from __future__ import annotations

import os
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any, Iterable, Optional
from urllib.parse import quote, urlencode

from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from .asoc_products import (
    ASOC_MATCH_CONFLICT,
    ASOC_MATCH_MAPPED,
    ASOC_MATCH_MISSING_TARGET,
    ASOC_MATCH_NO_REPOSITORIES,
    ASOC_MATCH_PARTIAL,
    CREATED_MAX,
    CREATED_MIN,
    _monitored_product_url,
    _monitored_repository_url,
    asoc_product_url,
    description_identity_urls,
    plausible_created,
    product_scope,
)
from .audit_service import AuditService, path_within
from .database import DatabaseService
from .display_time import DISPLAY_TIMEZONE
from .iac_asoc import normalize_repo_url
from .models import (
    ActivePipeline,
    AsocProduct,
    AsocProductRepositoryMapping,
    AuditCoverageStatus,
    Delivery,
    Execution,
    JobReport,
    PipelineStatus,
    Product,
    Repository,
)
from .url_safety import safe_external_http_url

REGISTRY_MODES = ("new", "total")
SCOPE_FILTERS = ("all", "internal", "external", "unclassified", "conflict")
STATUS_FILTERS = ("all", "synced", "attention", "removed")

STATUS_SYNCED = "synced"
STATUS_PARTIAL = "partial"
STATUS_CONFLICT = "conflict"
STATUS_GROUP_NOT_MONITORED = "group_not_monitored"
STATUS_NO_GITLAB_LINK = "no_gitlab_link"
STATUS_PENDING_SYNC = "pending_sync"
STATUS_REMOVED = "removed"

STATUS_LABELS = {
    STATUS_SYNCED: "Synced",
    STATUS_PARTIAL: "Partially synced",
    STATUS_CONFLICT: "Conflict",
    STATUS_GROUP_NOT_MONITORED: "Group not in ScanOps",
    STATUS_NO_GITLAB_LINK: "No GitLab link",
    STATUS_PENDING_SYNC: "Waiting for sync",
    STATUS_REMOVED: "Removed from ASOC",
}
STATUS_HINTS = {
    STATUS_SYNCED: "Every GitLab link on the ASOC card resolves to a monitored group.",
    STATUS_PARTIAL: "Some GitLab links on the ASOC card are not monitored by ScanOps yet.",
    STATUS_CONFLICT: "A GitLab link resolves to more than one product; fix the ASOC card.",
    STATUS_GROUP_NOT_MONITORED: "The GitLab group from the ASOC card has not reported to ScanOps yet.",
    STATUS_NO_GITLAB_LINK: "Add the GitLab group URL to the ASOC card description.",
    STATUS_PENDING_SYNC: "ScanOps has not matched this card yet.",
    STATUS_REMOVED: "The card is no longer present in ASOC.",
}
_ATTENTION_STATUSES = {
    STATUS_PARTIAL,
    STATUS_CONFLICT,
    STATUS_GROUP_NOT_MONITORED,
    STATUS_NO_GITLAB_LINK,
    STATUS_PENDING_SYNC,
}

REPO_SCANNED = "scanned"
REPO_NOT_SCANNED = "not_scanned"
REPO_INACTIVE = "inactive"
REPO_UNKNOWN = "unknown"
REPO_EXCLUDED = "excluded"
REPO_NOT_AUDITED = "not_audited"
REPO_STATUS_LABELS = {
    REPO_SCANNED: "Scanned",
    REPO_NOT_SCANNED: "Not scanned",
    REPO_INACTIVE: "No recent pipelines",
    REPO_UNKNOWN: "Audit failed",
    REPO_EXCLUDED: "Excluded",
    REPO_NOT_AUDITED: "Not audited",
}
_AUDIT_TO_REPO_STATUS = {
    AuditCoverageStatus.COVERED: REPO_SCANNED,
    AuditCoverageStatus.PIPELINE_NO_SEC: REPO_NOT_SCANNED,
    AuditCoverageStatus.NO_PIPELINES: REPO_INACTIVE,
    AuditCoverageStatus.INACCESSIBLE: REPO_UNKNOWN,
    AuditCoverageStatus.EXCLUDED: REPO_EXCLUDED,
}
_REPO_STATUS_ORDER = {
    REPO_NOT_SCANNED: 0,
    REPO_UNKNOWN: 1,
    REPO_NOT_AUDITED: 2,
    REPO_INACTIVE: 3,
    REPO_SCANNED: 4,
    REPO_EXCLUDED: 5,
}

PIPELINE_WINDOW = timedelta(days=30)
_TERMINAL_PIPELINE_STATUSES = (PipelineStatus.COMPLETED, PipelineStatus.FAILED)


def sync_grace_days() -> int:
    """Days a new ASOC card may wait for its GitLab group before it is overdue."""
    try:
        return max(int(os.getenv("PRODUCT_SYNC_GRACE_DAYS", "3")), 0)
    except ValueError:
        return 3


# ── Periods ──────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Quarter:
    year: int
    number: int

    @classmethod
    def parse(cls, value: str) -> "Quarter":
        raw = (value or "").strip().upper().replace(" ", "-")
        year_part, sep, quarter_part = raw.partition("-Q")
        if not sep or not year_part.isdigit() or quarter_part not in {"1", "2", "3", "4"}:
            raise ValueError("period must look like 2026-Q3")
        year = int(year_part)
        if not 2000 <= year <= 2100:
            raise ValueError("period year is out of range")
        return cls(year, int(quarter_part))

    @classmethod
    def containing(cls, value: datetime) -> "Quarter":
        local = _as_utc(value).astimezone(DISPLAY_TIMEZONE)
        return cls(local.year, (local.month - 1) // 3 + 1)

    @property
    def key(self) -> str:
        return f"{self.year}-Q{self.number}"

    @property
    def label(self) -> str:
        return f"{self.year} Q{self.number}"

    @property
    def start(self) -> datetime:
        month = (self.number - 1) * 3 + 1
        return datetime(self.year, month, 1, tzinfo=DISPLAY_TIMEZONE).astimezone(timezone.utc)

    @property
    def end(self) -> datetime:
        return self.next().start

    def next(self) -> "Quarter":
        if self.number == 4:
            return Quarter(self.year + 1, 1)
        return Quarter(self.year, self.number + 1)


def quarter_label(value: Optional[datetime]) -> Optional[str]:
    return Quarter.containing(value).label if value else None


@dataclass(frozen=True)
class RegistryQuery:
    period: Optional[Quarter] = None
    mode: str = "new"
    scope: str = "all"
    status: str = "all"
    q: str = ""

    @classmethod
    def from_params(
        cls,
        *,
        period: Optional[str] = None,
        mode: Optional[str] = None,
        scope: Optional[str] = None,
        status: Optional[str] = None,
        q: Optional[str] = None,
    ) -> "RegistryQuery":
        """Validate URL parameters; empty values mean "no filter"."""
        parsed_period = Quarter.parse(period) if period else None
        mode_value = (mode or "new").strip().lower()
        if mode_value not in REGISTRY_MODES:
            raise ValueError("mode must be new or total")
        scope_value = (scope or "all").strip().lower()
        if scope_value not in SCOPE_FILTERS:
            raise ValueError("scope is not supported")
        status_value = (status or "all").strip().lower()
        if status_value not in STATUS_FILTERS:
            raise ValueError("status is not supported")
        return cls(parsed_period, mode_value, scope_value, status_value, (q or "").strip()[:200])

    def as_params(self) -> dict[str, str]:
        params: dict[str, str] = {}
        if self.period:
            params["period"] = self.period.key
            params["mode"] = self.mode
        if self.scope != "all":
            params["scope"] = self.scope
        if self.status != "all":
            params["status"] = self.status
        if self.q:
            params["q"] = self.q
        return params


def _as_utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _iso(value: Optional[datetime]) -> Optional[str]:
    return _as_utc(value).isoformat() if value else None


def card_in_period(card: AsocProduct, query: RegistryQuery) -> bool:
    """Apply the quarterly slice to one ASOC card.

    ``new`` keeps cards created in ASOC during the quarter, including cards
    removed later. ``total`` keeps every card that existed at the end of the
    quarter: created before the end and not removed before it.
    """
    if query.period is None:
        return bool(card.is_active)
    created = plausible_created(card.asoc_created)
    if created is None:
        return False
    start, end = query.period.start, query.period.end
    if query.mode == "new":
        return start <= created < end
    if created >= end:
        return False
    if card.is_active:
        return True
    return card.removed_at is not None and _as_utc(card.removed_at) >= end


def card_status(card: AsocProduct) -> str:
    if not card.is_active:
        return STATUS_REMOVED
    return {
        ASOC_MATCH_MAPPED: STATUS_SYNCED,
        ASOC_MATCH_PARTIAL: STATUS_PARTIAL,
        ASOC_MATCH_CONFLICT: STATUS_CONFLICT,
        ASOC_MATCH_MISSING_TARGET: STATUS_GROUP_NOT_MONITORED,
        ASOC_MATCH_NO_REPOSITORIES: STATUS_NO_GITLAB_LINK,
    }.get(card.mapping_state or "", STATUS_PENDING_SYNC)


def status_group(status: str) -> str:
    if status == STATUS_SYNCED:
        return "synced"
    if status == STATUS_REMOVED:
        return "removed"
    return "attention"


def available_periods(db: Session, now: Optional[datetime] = None) -> list[dict[str, str]]:
    """Quarters from the first ASOC connection to the current quarter, newest first."""
    earliest = (
        db.query(func.min(AsocProduct.asoc_created))
        .filter(AsocProduct.asoc_created >= CREATED_MIN, AsocProduct.asoc_created < CREATED_MAX)
        .scalar()
    )
    current = Quarter.containing(now or datetime.now(timezone.utc))
    periods: list[Quarter] = []
    if earliest is not None:
        quarter = Quarter.containing(earliest)
        while (quarter.year, quarter.number) <= (current.year, current.number):
            periods.append(quarter)
            quarter = quarter.next()
    if not periods:
        periods = [current]
    return [{"key": item.key, "label": item.label} for item in reversed(periods)]


# ── Evidence collection ─────────────────────────────────────────────────


def card_href(asoc_id: str) -> str:
    """Drill-down URL of an ASOC card; the id is ASOC-controlled text."""
    return f"/targets/products/asoc/{quote(asoc_id, safe='')}"


def _group_path(product: Product) -> str:
    return (product.path or "").strip().strip("/")


def _group_url(product: Product) -> Optional[str]:
    path = _group_path(product)
    if path.startswith(("http://", "https://")):
        return safe_external_http_url(path)
    base = (product.gitlab_base_url or "").strip().rstrip("/")
    return safe_external_http_url(f"{base}/{path}" if base and path else "")


def _group_identity_url(product: Product) -> str:
    path = _group_path(product)
    if path.startswith(("http://", "https://")):
        return normalize_repo_url(path)
    base = (product.gitlab_base_url or "").strip().rstrip("/")
    return normalize_repo_url(f"{base}/{path}") if base and path else ""


def _repository_url(product: Optional[Product], path: str) -> Optional[str]:
    if path.startswith(("http://", "https://")):
        return safe_external_http_url(path)
    base = ((product.gitlab_base_url if product else "") or "").strip().rstrip("/")
    return safe_external_http_url(f"{base}/{path.strip('/')}" if base and path else "")


def _pipeline_terminal_at():
    return func.coalesce(
        ActivePipeline.completed_at,
        ActivePipeline.finished_at,
        ActivePipeline.updated_at,
        ActivePipeline.created_at,
    )


class _Evidence:
    """Batch-loaded ScanOps evidence shared by every row of one request."""

    def __init__(self, db: Session, now: datetime):
        self.db = db
        self.now = now
        self.products: dict[int, Product] = {
            product.id: product for product in db.query(Product).all()
        }
        self.repositories_by_product: dict[int, list[Repository]] = defaultdict(list)
        for repository in db.query(Repository).order_by(Repository.path.asc()).all():
            self.repositories_by_product[repository.product_id].append(repository)
        self.targets_by_card: dict[str, list[int]] = defaultdict(list)
        self.links_by_card: dict[str, dict[str, int]] = defaultdict(dict)
        for mapping in db.query(AsocProductRepositoryMapping).order_by(
            AsocProductRepositoryMapping.product_id.asc()
        ):
            if mapping.product_id not in self.targets_by_card[mapping.asoc_id]:
                self.targets_by_card[mapping.asoc_id].append(mapping.product_id)
            self.links_by_card[mapping.asoc_id][mapping.description_url] = mapping.product_id
        self.audit_by_product = AuditService(db).latest_repositories_by_product()
        self._pipeline_stats: Optional[dict[int, dict[str, Any]]] = None

    def targets_for(self, card: AsocProduct) -> list[int]:
        targets = [
            product_id
            for product_id in self.targets_by_card.get(card.asoc_id, [])
            if product_id in self.products
        ]
        if not targets and card.monitor_product_id in self.products and card.mapping_state != ASOC_MATCH_CONFLICT:
            targets = [card.monitor_product_id]
        return targets

    @property
    def pipeline_stats(self) -> dict[int, dict[str, Any]]:
        """Per-repository pipeline outcomes: 30-day counts and the latest result."""
        if self._pipeline_stats is not None:
            return self._pipeline_stats
        stats: dict[int, dict[str, Any]] = defaultdict(
            lambda: {"completed": 0, "failed": 0, "last_status": None, "last_at": None, "last_url": None, "last_success_at": None}
        )
        terminal_at = _pipeline_terminal_at()
        window_start = self.now - PIPELINE_WINDOW
        for repository_id, status, count in (
            self.db.query(ActivePipeline.repository_id, ActivePipeline.status, func.count(ActivePipeline.id))
            .filter(ActivePipeline.status.in_(_TERMINAL_PIPELINE_STATUSES))
            .filter(terminal_at >= window_start)
            .group_by(ActivePipeline.repository_id, ActivePipeline.status)
        ):
            key = "completed" if _status_value(status) == PipelineStatus.COMPLETED.value else "failed"
            stats[repository_id][key] += int(count)

        latest = (
            self.db.query(
                ActivePipeline.repository_id.label("repository_id"),
                func.max(terminal_at).label("terminal_at"),
            )
            .filter(ActivePipeline.status.in_(_TERMINAL_PIPELINE_STATUSES))
            .group_by(ActivePipeline.repository_id)
            .subquery()
        )
        for pipeline in (
            self.db.query(ActivePipeline)
            .join(
                latest,
                and_(
                    ActivePipeline.repository_id == latest.c.repository_id,
                    terminal_at == latest.c.terminal_at,
                ),
            )
            .filter(ActivePipeline.status.in_(_TERMINAL_PIPELINE_STATUSES))
            .order_by(ActivePipeline.id.asc())
        ):
            row = stats[pipeline.repository_id]
            row["last_status"] = _status_value(pipeline.status)
            row["last_at"] = _pipeline_terminal_value(pipeline)
            row["last_url"] = safe_external_http_url(pipeline.web_url)

        for repository_id, value in (
            self.db.query(ActivePipeline.repository_id, func.max(terminal_at))
            .filter(ActivePipeline.status == PipelineStatus.COMPLETED)
            .group_by(ActivePipeline.repository_id)
        ):
            stats[repository_id]["last_success_at"] = value
        self._pipeline_stats = stats
        return stats


def _status_value(value: Any) -> str:
    return value.value if hasattr(value, "value") else str(value)


def _pipeline_terminal_value(pipeline: ActivePipeline) -> Optional[datetime]:
    return pipeline.completed_at or pipeline.finished_at or pipeline.updated_at or pipeline.created_at


def _repository_inventory(evidence: _Evidence, group_ids: Iterable[int]) -> list[dict[str, Any]]:
    """Union of audited GitLab repositories and ScanOps repositories for groups."""
    inventory: dict[str, dict[str, Any]] = {}
    for group_id in group_ids:
        product = evidence.products.get(group_id)
        group_path = AuditService.product_group_path(product.path if product else None)
        for snapshot, audit_repo in evidence.audit_by_product.get(group_id, []):
            path = (audit_repo.path_with_namespace or "").strip("/")
            key = path.lower()
            # The audit may attribute a project by its top-level namespace only;
            # the registry lists repositories that are inside the product group.
            if group_path and not path_within(key, group_path):
                continue
            status = _AUDIT_TO_REPO_STATUS.get(snapshot.status, REPO_UNKNOWN)
            if audit_repo.excluded or audit_repo.archived:
                status = REPO_EXCLUDED
            inventory[key] = {
                "path": path,
                "url": safe_external_http_url(audit_repo.web_url) or _repository_url(product, path),
                "group_id": group_id,
                "scanops_repository_id": None,
                "status": status,
                "audit_reason": snapshot.reason,
                "audited_at": snapshot.scanned_at,
                "last_pipeline_at": snapshot.last_pipeline_at,
            }
        for repository in evidence.repositories_by_product.get(group_id, []):
            path = (repository.path or "").strip("/")
            key = path.lower()
            entry = inventory.get(key)
            if entry is None:
                entry = {
                    "path": path,
                    "url": _repository_url(product, path),
                    "group_id": group_id,
                    "scanops_repository_id": None,
                    "status": REPO_NOT_AUDITED,
                    "audit_reason": None,
                    "audited_at": None,
                    "last_pipeline_at": None,
                }
                inventory[key] = entry
            entry["scanops_repository_id"] = repository.id
    return sorted(
        inventory.values(),
        key=lambda item: (_REPO_STATUS_ORDER.get(item["status"], 9), item["path"].lower()),
    )


def _repository_summary(inventory: list[dict[str, Any]], evidence: _Evidence) -> dict[str, Any]:
    counts = Counter(item["status"] for item in inventory)
    audited = [item["audited_at"] for item in inventory if item["audited_at"]]
    pipeline_rows = [
        evidence.pipeline_stats[item["scanops_repository_id"]]
        for item in inventory
        if item["scanops_repository_id"] is not None
        and item["scanops_repository_id"] in evidence.pipeline_stats
    ]
    completed = sum(row["completed"] for row in pipeline_rows)
    failed = sum(row["failed"] for row in pipeline_rows)
    with_pipelines = [row for row in pipeline_rows if row["last_status"]]
    last_success = max((row["last_success_at"] for row in pipeline_rows if row["last_success_at"]), default=None)
    return {
        "repositories": {
            # Excluded (archived/empty/opted-out) projects are not expected to scan.
            "total": sum(count for status, count in counts.items() if status != REPO_EXCLUDED),
            "scanned": counts.get(REPO_SCANNED, 0),
            "not_scanned": counts.get(REPO_NOT_SCANNED, 0),
            "inactive": counts.get(REPO_INACTIVE, 0),
            "unknown": counts.get(REPO_UNKNOWN, 0),
            "not_audited": counts.get(REPO_NOT_AUDITED, 0),
            "excluded": counts.get(REPO_EXCLUDED, 0),
            "monitored": sum(1 for item in inventory if item["scanops_repository_id"] is not None),
            "audited": bool(audited),
            "audited_at": _iso(max(audited)) if audited else None,
        },
        "pipelines": {
            "completed": completed,
            "failed": failed,
            "success_rate": round(completed / (completed + failed) * 100) if completed + failed else None,
            "repos_with_pipelines": len(with_pipelines),
            "repos_green": sum(1 for row in with_pipelines if row["last_status"] == PipelineStatus.COMPLETED.value),
            "last_success_at": _iso(last_success),
            "window_days": PIPELINE_WINDOW.days,
        },
    }


def _group_payload(evidence: _Evidence, product_id: int) -> dict[str, Any]:
    product = evidence.products[product_id]
    return {
        "id": product.id,
        "name": product.name,
        "path": _group_path(product),
        "url": _group_url(product),
        "href": f"/targets/products/{product.id}",
        "operator_href": f"/targets/products/{product.id}/operator",
        "pipelines_href": f"/pipelines?product={product.id}",
    }


def _card_row(
    card: AsocProduct,
    evidence: _Evidence,
    *,
    asoc_base_url: Optional[str],
    grace_days: int,
) -> dict[str, Any]:
    status = card_status(card)
    group_ids = evidence.targets_for(card) if status != STATUS_CONFLICT else []
    inventory = _repository_inventory(evidence, group_ids)
    metrics = _repository_summary(inventory, evidence)
    connected_at = plausible_created(card.asoc_created)
    # Waiting measures how long a card has sat in ASOC without ScanOps linking
    # it. A card with a linked group, or a recorded link date, is linked: its
    # attention status is a degradation, not an onboarding delay. Cards linked
    # before date tracking have groups but no date.
    currently_linked = bool(evidence.targets_for(card))
    pre_tracking = bool(card.connected_before_tracking) or currently_linked
    ever_linked = card.first_connected_at is not None or pre_tracking
    waiting_days = None
    if status in _ATTENTION_STATUSES and connected_at is not None and not ever_linked:
        waiting_days = max((evidence.now - connected_at).days, 0)
    gitlab_links = [
        url
        for url in description_identity_urls(card.description)
        if "gitlab" in url.split("/", 3)[2].casefold() or url in evidence.links_by_card.get(card.asoc_id, {})
    ]
    return {
        "key": f"asoc:{card.asoc_id}",
        "asoc_id": card.asoc_id,
        "bi_id": card.bi_id,
        "name": card.title or card.name,
        "asoc_name": card.name,
        "asoc_url": safe_external_http_url(
            asoc_product_url(card.asoc_id, card.slug, base_url=asoc_base_url)
        ),
        "scope": product_scope(card.description),
        "status": status,
        "status_group": status_group(status),
        "status_label": STATUS_LABELS[status],
        "status_hint": STATUS_HINTS[status],
        "connected_at": _iso(connected_at),
        "connected_quarter": quarter_label(connected_at),
        "linked_at": _iso(card.first_connected_at),
        "linked_before_tracking": card.first_connected_at is None and pre_tracking,
        "waiting_days": waiting_days,
        "overdue": waiting_days is not None and waiting_days > grace_days,
        "is_active": bool(card.is_active),
        "removed_at": _iso(card.removed_at),
        "links_total": len(gitlab_links),
        "links_matched": len(evidence.links_by_card.get(card.asoc_id, {})),
        "groups": [_group_payload(evidence, product_id) for product_id in group_ids],
        **metrics,
        "href": card_href(card.asoc_id),
    }


def _group_row(evidence: _Evidence, product_id: int) -> dict[str, Any]:
    inventory = _repository_inventory(evidence, [product_id])
    return {
        "key": f"group:{product_id}",
        **_group_payload(evidence, product_id),
        **_repository_summary(inventory, evidence),
    }


def _matches_search(row: dict[str, Any], needle: str) -> bool:
    if not needle:
        return True
    haystack = " ".join(
        [
            str(row.get("name") or ""),
            str(row.get("asoc_name") or ""),
            str(row.get("bi_id") or ""),
            f"BI_{row.get('bi_id')}" if row.get("bi_id") else "",
            str(row.get("asoc_id") or ""),
            str(row.get("path") or ""),
            *(f"{group['name']} {group['path']}" for group in row.get("groups", [])),
        ]
    ).casefold()
    return needle.casefold() in haystack


def _summary(rows: list[dict[str, Any]], query: RegistryQuery, evidence: _Evidence) -> dict[str, Any]:
    groups = Counter(row["status_group"] for row in rows)
    # Several cards (e.g. a removed card and its replacement) can link the same
    # GitLab group; count each repository once across the visible slice.
    group_ids = sorted({group["id"] for row in rows for group in row["groups"]})
    totals = _repository_summary(_repository_inventory(evidence, group_ids), evidence)
    repositories = {
        key: totals["repositories"][key]
        for key in ("total", "scanned", "not_scanned", "inactive", "unknown", "not_audited")
    }
    pipelines = totals["pipelines"]
    terminal = pipelines["completed"] + pipelines["failed"]
    if query.period is None:
        new_in_period = None
    elif query.mode == "new":
        new_in_period = len(rows)
    else:
        new_in_period = sum(
            1 for row in rows if row["connected_quarter"] == query.period.label
        )
    return {
        "total": len(rows),
        "synced": groups.get("synced", 0),
        "attention": groups.get("attention", 0),
        "removed": groups.get("removed", 0),
        "overdue": sum(1 for row in rows if row["overdue"]),
        "new_in_period": new_in_period,
        "repositories": repositories,
        "pipelines": {
            "completed": pipelines["completed"],
            "failed": pipelines["failed"],
            "success_rate": round(pipelines["completed"] / terminal * 100) if terminal else None,
        },
    }


def build_registry(
    db: Session,
    query: RegistryQuery,
    *,
    asoc_base_url: Optional[str] = None,
    now: Optional[datetime] = None,
) -> dict[str, Any]:
    current = _as_utc(now or datetime.now(timezone.utc))
    evidence = _Evidence(db, current)
    grace_days = sync_grace_days()
    cards = db.query(AsocProduct).order_by(AsocProduct.asoc_id.asc()).all()

    sliced = [card for card in cards if card_in_period(card, query)]
    rows = [
        _card_row(card, evidence, asoc_base_url=asoc_base_url, grace_days=grace_days)
        for card in sliced
    ]
    rows = [
        row
        for row in rows
        if (query.scope == "all" or row["scope"] == query.scope)
        and _matches_search(row, query.q)
    ]
    rows.sort(key=lambda row: (str(row["name"]).casefold(), row["asoc_id"]))
    summary = _summary(rows, query, evidence)
    visible = [row for row in rows if query.status == "all" or row["status_group"] == query.status]

    unregistered: list[dict[str, Any]] = []
    if query.period is None:
        claimed = {product_id for card in cards if card.is_active for product_id in evidence.targets_for(card)}
        for product_id, product in evidence.products.items():
            if product_id in claimed:
                continue
            # Conflicted groups stay listed: their page explains which cards
            # compete for them, and nothing else in Products links there.
            row = {**_group_row(evidence, product_id), "asoc_conflict": product.asoc_sync_status == ASOC_MATCH_CONFLICT}
            if _matches_search(row, query.q):
                unregistered.append(row)
        unregistered.sort(key=lambda row: str(row["name"]).casefold())

    # Cards without a usable creation date cannot be placed in any quarter;
    # count the ones the current audience/search filters would otherwise show.
    undated = sum(
        1
        for card in cards
        if plausible_created(card.asoc_created) is None
        and (card.is_active or query.period is not None)
        and (query.scope == "all" or product_scope(card.description) == query.scope)
        and _matches_search(
            _card_row(card, evidence, asoc_base_url=asoc_base_url, grace_days=grace_days),
            query.q,
        )
    )
    return {
        "query": {
            "period": query.period.key if query.period else None,
            "period_label": query.period.label if query.period else None,
            "mode": query.mode,
            "scope": query.scope,
            "status": query.status,
            "q": query.q,
        },
        "periods": available_periods(db, current),
        "rows": visible,
        "summary": {**summary, "undated": undated, "grace_days": grace_days},
        "unregistered_groups": unregistered,
        "generated_at": current.isoformat(),
    }


# ── Product drill-down ──────────────────────────────────────────────────


def _open_failed_jobs(db: Session, repository_ids: list[int]) -> dict[int, int]:
    if not repository_ids:
        return {}
    return {
        int(repository_id): int(count)
        for repository_id, count in DatabaseService(db)
        ._active_failed_jobs_query()
        .filter(JobReport.repository_id.in_(repository_ids))
        .with_entities(JobReport.repository_id, func.count(JobReport.id))
        .group_by(JobReport.repository_id)
    }


def asoc_delivery_filter():
    """Select only canonical ASOC destinations, never look-alikes or archives."""
    return or_(Delivery.destination == "asoc", Delivery.destination.like("asoc:%"))


def _asoc_deliveries(db: Session, group_ids: list[int], repository_ids: list[int]) -> dict[str, Any]:
    if not group_ids and not repository_ids:
        return {"summary": {"total": 0, "confirmed": 0, "failed": 0, "pending": 0}, "rows": []}
    scope = []
    if group_ids:
        scope.append(Execution.product_id.in_(group_ids))
    if repository_ids:
        scope.append(Execution.repository_id.in_(repository_ids))
    base = (
        db.query(Delivery, Execution)
        .join(Execution, Delivery.execution_id == Execution.id)
        .filter(asoc_delivery_filter())
        .filter(or_(*scope))
    )
    counts = Counter(
        {
            status: int(count)
            for status, count in base.with_entities(Delivery.status, func.count(Delivery.id)).group_by(Delivery.status)
        }
    )
    rows = []
    for delivery, execution in base.order_by(
        func.coalesce(Delivery.attempted_at, Delivery.created_at).desc(), Delivery.id.desc()
    ).limit(20):
        rows.append(
            {
                "id": delivery.id,
                "execution_id": execution.id,
                "source": execution.source,
                "status": delivery.status,
                "failure_reason": delivery.failure_reason,
                "attempted_at": _iso(delivery.attempted_at),
                "confirmed_at": _iso(delivery.confirmed_at),
                "repository_id": execution.repository_id,
            }
        )
    return {
        "summary": {
            "total": sum(counts.values()),
            "confirmed": counts.get("confirmed", 0),
            "failed": counts.get("failed", 0),
            "pending": counts.get("requested", 0) + counts.get("running", 0),
        },
        "rows": rows,
    }


def _detail_repositories(db: Session, evidence: _Evidence, inventory: list[dict[str, Any]]) -> list[dict[str, Any]]:
    repository_ids = [item["scanops_repository_id"] for item in inventory if item["scanops_repository_id"]]
    failed_jobs = _open_failed_jobs(db, repository_ids)
    rows = []
    for item in inventory:
        repository_id = item["scanops_repository_id"]
        pipeline = evidence.pipeline_stats.get(repository_id) if repository_id else None
        completed = pipeline["completed"] if pipeline else 0
        failed = pipeline["failed"] if pipeline else 0
        group = evidence.products.get(item["group_id"])
        rows.append(
            {
                "path": item["path"],
                "url": item["url"],
                "group_id": item["group_id"],
                "group_name": group.name if group else None,
                "status": item["status"],
                "status_label": REPO_STATUS_LABELS[item["status"]],
                "audit_reason": item["audit_reason"],
                "audited_at": _iso(item["audited_at"]),
                "last_activity_at": _iso(item["last_pipeline_at"]),
                "scanops_repository_id": repository_id,
                "href": f"/targets/repositories/{repository_id}" if repository_id else None,
                "jobs_href": "/logs?" + urlencode({"repo": item["path"], "status": "failed"}) if repository_id else None,
                "open_failed_jobs": failed_jobs.get(repository_id, 0) if repository_id else 0,
                "pipelines": {
                    "completed": completed,
                    "failed": failed,
                    "success_rate": round(completed / (completed + failed) * 100) if completed + failed else None,
                    "last_status": pipeline["last_status"] if pipeline else None,
                    "last_at": _iso(pipeline["last_at"]) if pipeline else None,
                    "last_url": pipeline["last_url"] if pipeline else None,
                },
            }
        )
    return rows


def _description_links(card: AsocProduct, evidence: _Evidence) -> list[dict[str, Any]]:
    mapped = evidence.links_by_card.get(card.asoc_id, {})
    links = []
    for url in description_identity_urls(card.description):
        host = url.split("/", 3)[2].casefold() if "://" in url else ""
        is_gitlab = "gitlab" in host or url in mapped
        product_id = mapped.get(url)
        if not is_gitlab:
            state = "reference"
        elif product_id in evidence.products:
            state = "matched"
        elif card.mapping_state == ASOC_MATCH_CONFLICT:
            state = "ambiguous"
        else:
            state = "not_monitored"
        links.append(
            {
                "url": safe_external_http_url(url),
                "kind": "gitlab" if is_gitlab else "reference",
                "state": state,
                "group": _group_payload(evidence, product_id) if state == "matched" else None,
            }
        )
    return links


def product_detail(
    db: Session,
    *,
    asoc_id: Optional[str] = None,
    product_id: Optional[int] = None,
    asoc_base_url: Optional[str] = None,
    now: Optional[datetime] = None,
) -> Optional[dict[str, Any]]:
    """Drill-down for one ASOC card, or for a monitored group without a card."""
    current = _as_utc(now or datetime.now(timezone.utc))
    evidence = _Evidence(db, current)
    card: Optional[AsocProduct] = None
    focus_group_id: Optional[int] = None
    if asoc_id is not None:
        card = db.query(AsocProduct).filter(AsocProduct.asoc_id == asoc_id).first()
        if card is None:
            return None
    else:
        if product_id not in evidence.products:
            return None
        focus_group_id = product_id
        owners = [
            candidate
            for candidate in db.query(AsocProduct).filter(AsocProduct.is_active.is_(True))
            if product_id in evidence.targets_for(candidate)
        ]
        if len(owners) == 1:
            card = owners[0]

    if card is not None:
        row = _card_row(card, evidence, asoc_base_url=asoc_base_url, grace_days=sync_grace_days())
        group_ids = [group["id"] for group in row["groups"]]
        inventory = _repository_inventory(evidence, group_ids)
        kind = "asoc"
        links = _description_links(card, evidence)
        claimed_by: list[dict[str, Any]] = []
    else:
        row = _group_row(evidence, product_id)
        group_ids = [product_id]
        inventory = _repository_inventory(evidence, group_ids)
        kind = "group"
        links = []
        group_url = _group_identity_url(evidence.products[product_id])
        repo_urls = {
            normalize_repo_url(item["url"]) for item in inventory if item["url"]
        }
        # Build the URLs exactly as the sync does: it resolves the group URL
        # and repositories ScanOps already knows, nothing else.
        group = evidence.products[product_id]
        syncable_urls = {
            url
            for url in (
                _monitored_product_url(group),
                *(
                    _monitored_repository_url(group, repository)
                    for repository in evidence.repositories_by_product.get(product_id, [])
                ),
            )
            if url
        }
        group_identity = _monitored_product_url(group)
        # The sync refuses to pick an owner when a group URL matches several
        # ScanOps products.
        group_url_owners = sum(
            1
            for other in evidence.products.values()
            if group_identity and _monitored_product_url(other) == group_identity
        )
        # A shared group URL never resolves, so it cannot make a card a rival.
        competing_urls = syncable_urls - {group_identity} if group_url_owners > 1 else syncable_urls
        claimed_by = []
        competes: dict[str, bool] = {}
        links_group: dict[str, bool] = {}
        for candidate in db.query(AsocProduct).filter(AsocProduct.is_active.is_(True)).order_by(AsocProduct.asoc_id):
            linked = set(description_identity_urls(candidate.description))
            if not ({group_url, *repo_urls} & linked):
                continue
            competes[candidate.asoc_id] = bool(competing_urls & linked)
            links_group[candidate.asoc_id] = bool(group_identity and group_identity in linked)
            if candidate.mapping_state == ASOC_MATCH_CONFLICT or candidate in owners:
                # Competing cards stay in conflict and own nothing.
                state = "conflict"
            elif syncable_urls & linked:
                state = "pending_sync"
            else:
                # The card links this group's projects in a form the sync cannot
                # resolve (e.g. audit-only repositories): it will not connect.
                state = "unresolved"
            claimed_by.append({
                "asoc_id": candidate.asoc_id,
                "name": candidate.title or candidate.name,
                "bi_id": candidate.bi_id,
                "href": card_href(candidate.asoc_id),
                "state": state,
                "status_label": STATUS_LABELS[card_status(candidate)],
            })
        # A waiting card will not win the group if another card resolves to it
        # too (already conflicted or also waiting), or if it links a group URL
        # that several ScanOps products share.
        for claim in claimed_by:
            if claim["state"] != "pending_sync":
                continue
            rivals = any(
                other is not claim and competes[other["asoc_id"]] and other["state"] in {"pending_sync", "conflict"}
                for other in claimed_by
            )
            if rivals or (links_group[claim["asoc_id"]] and group_url_owners > 1):
                claim["state"] = "conflict"

    repositories = _detail_repositories(db, evidence, inventory)
    repository_ids = [item["scanops_repository_id"] for item in repositories if item["scanops_repository_id"]]
    return {
        "kind": kind,
        "focus_group_id": focus_group_id,
        "product": row,
        "links": links,
        "claimed_by": claimed_by,
        "repositories": repositories,
        "open_failed_jobs": sum(item["open_failed_jobs"] for item in repositories),
        "deliveries": _asoc_deliveries(db, group_ids, repository_ids),
        "generated_at": current.isoformat(),
    }


# ── Export ──────────────────────────────────────────────────────────────

EXPORT_COLUMNS = [
    ("BI_ID", 12),
    ("Product (ASOC)", 42),
    ("ASOC link", 44),
    ("GitLab groups", 40),
    ("GitLab group links", 48),
    ("Repositories in groups", 14),
    ("Scanned repositories", 14),
    ("Not scanned (active)", 14),
    ("No recent pipelines", 14),
    ("Pipeline success, 30d", 14),
    ("Repositories with green last pipeline", 18),
    ("Sync status", 26),
    ("Connected (ASOC created)", 16),
    ("Quarter", 10),
    ("Audience", 14),
    ("Removed from ASOC", 16),
    ("Last successful pipeline", 18),
]


def _local_date(value: Optional[str]) -> str:
    if not value:
        return ""
    return datetime.fromisoformat(value).astimezone(DISPLAY_TIMEZONE).strftime("%Y-%m-%d")


def _append_text_row(sheet: Any, values: list[Any]) -> None:
    """Append a row whose strings stay inert text in every spreadsheet app."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    sheet.append([
        ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
        for value in values
    ])
    for cell in sheet[sheet.max_row]:
        # ASOC-controlled strings must never become spreadsheet formulas.
        if isinstance(cell.value, str):
            cell.data_type = "s"


def export_xlsx(registry: dict[str, Any]) -> bytes:
    """One workbook per registry view: Products plus the slice summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append([name for name, _width in EXPORT_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in registry["rows"]:
        repositories = row["repositories"]
        pipelines = row["pipelines"]
        values = [
            f"BI_{row['bi_id']}" if row["bi_id"] else "",
            row["name"],
            row["asoc_url"] or "",
            "; ".join(group["path"] or group["name"] for group in row["groups"]),
            "; ".join(group["url"] or "" for group in row["groups"]),
            repositories["total"],
            repositories["scanned"],
            repositories["not_scanned"],
            repositories["inactive"],
            pipelines["success_rate"] if pipelines["success_rate"] is not None else "",
            f"{pipelines['repos_green']}/{pipelines['repos_with_pipelines']}" if pipelines["repos_with_pipelines"] else "",
            row["status_label"],
            _local_date(row["connected_at"]),
            row["connected_quarter"] or "",
            row["scope"],
            _local_date(row["removed_at"]),
            _local_date(pipelines["last_success_at"]),
        ]
        _append_text_row(sheet, values)
        if row["asoc_url"]:
            sheet.cell(row=sheet.max_row, column=3).hyperlink = row["asoc_url"]
    for index, (_name, width) in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    summary_sheet = workbook.create_sheet("Summary")
    query = registry["query"]
    summary = registry["summary"]
    if query["period"]:
        period_text = f"{query['period_label']} · {'new in quarter' if query['mode'] == 'new' else 'connected at quarter end'}"
    else:
        period_text = "Current registry"
    lines = [
        ("Period", period_text),
        ("Audience", query["scope"]),
        ("Status filter", query["status"]),
        ("Search", query["q"]),
        ("Generated", datetime.fromisoformat(registry["generated_at"]).astimezone(DISPLAY_TIMEZONE).strftime("%Y-%m-%d %H:%M %Z")),
        ("Products", summary["total"]),
        ("Synced", summary["synced"]),
        ("Need attention", summary["attention"]),
        ("Removed from ASOC", summary["removed"]),
        ("Repositories in groups", summary["repositories"].get("total", 0)),
        ("Scanned repositories", summary["repositories"].get("scanned", 0)),
        ("Not scanned (active)", summary["repositories"].get("not_scanned", 0)),
        ("No recent pipelines", summary["repositories"].get("inactive", 0)),
        ("Pipeline success, 30d", summary["pipelines"]["success_rate"] if summary["pipelines"]["success_rate"] is not None else ""),
        ("ASOC cards without a creation date", summary["undated"]),
    ]
    if summary["new_in_period"] is not None:
        lines.insert(6, ("New in quarter", summary["new_in_period"]))
    for label, value in lines:
        _append_text_row(summary_sheet, [label, value])
    summary_sheet.column_dimensions["A"].width = 26
    summary_sheet.column_dimensions["B"].width = 48
    for cell in summary_sheet["A"]:
        cell.font = Font(bold=True)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(registry: dict[str, Any]) -> str:
    query = registry["query"]
    if query["period"]:
        suffix = f"{query['period']}-{'new' if query['mode'] == 'new' else 'total'}"
    else:
        suffix = "current"
    return f"scanops-products-{suffix}.xlsx"
